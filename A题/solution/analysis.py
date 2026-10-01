# -*- coding: utf-8 -*-
"""模型验证(解析解对比) + 灵敏度分析，供论文引用"""
import numpy as np
from scipy.special import j0, j1
from solver import (simulate, simulate_until, make_air_functions, make_R_function,
                    props_p1, props_p23, props_p4)

DATA = r'C:/Users/hp/Downloads/CUMCM2026Problems/A题/附件'
import openpyxl

def load_air():
    wb = openpyxl.load_workbook(DATA + r'/附件1.xlsx', data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))[1:]
    return (np.array([r[0] for r in rows], float),
            np.array([r[1] for r in rows], float),
            np.array([r[2] for r in rows], float))

def load_radius():
    wb = openpyxl.load_workbook(DATA + r'/附件2.xlsx', data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))[1:]
    return (np.array([r[0] for r in rows], float),
            np.array([r[1] for r in rows], float))

t_air, T_air, C_air = load_air()
t_rad, R_rad = load_radius()
T_air_fn, C_air_fn = make_air_functions(t_air, T_air, C_air, 50.0, 0.05)
R_fn = make_R_function(t_rad, R_rad)

R0 = 0.02; T0 = 28.0; C0 = 2.55; h = 25.0; hm = 8e-7

# ================= 验证: 解析解 =================
def analytic_cyl(r, t, R, U0, Uinf, alpha, Bi, nmax=300):
    r = np.atleast_1d(np.asarray(r, float))
    lam = []
    x = 0.05
    while len(lam) < nmax and x < 1e4:
        a, b = x, x + 0.05
        fa = a*j1(a) - Bi*j0(a)
        # 粗扫根
        xx = np.linspace(a, b, 9)
        f = xx*j1(xx) - Bi*j0(xx)
        for k in range(len(xx)-1):
            if f[k]*f[k+1] < 0:
                lo, hi = xx[k], xx[k+1]
                for _ in range(60):
                    m = 0.5*(lo+hi)
                    if (lo*j1(lo)-Bi*j0(lo))*(m*j1(m)-Bi*j0(m)) <= 0:
                        hi = m
                    else:
                        lo = m
                rt = 0.5*(lo+hi)
                if not lam or rt - lam[-1] > 1e-5:
                    lam.append(rt)
        x += 0.05
    lam = np.array(lam)
    res = np.zeros((len(r), len(t)))
    for k, tt in enumerate(t):
        Fo = alpha*tt/R**2
        for jj, rj in enumerate(r):
            s = 0.0
            for ln in lam:
                An = 2.0*j1(ln)/(ln*(j0(ln)**2 + j1(ln)**2))
                s += An*np.exp(-ln**2*Fo)*j0(ln*rj/R)
            res[jj, k] = Uinf + (U0 - Uinf)*s
    return res

print('='*60)
print('验证1: 热传导(常数物性)')
rho, cp, k = 820.0, 2600.0, 0.36
alpha = k/(rho*cp); Bi = h*R0/k
times = [100, 300, 600, 1200, 1800]
simT = simulate(lambda t: 50.0, lambda t: 0.0, lambda t: R0,
                lambda C, T: (rho, cp, k, 0.0), 400, T0, 0.0, 1800.0, 0.5,
                h, 1e-9, save_times=times)
for rpt in [0.0, 0.01, 0.02]:
    idx = int(round(rpt/R0*400))
    ana = analytic_cyl([rpt], times, R0, T0, 50.0, alpha, Bi)
    errs = [abs(simT['T'][j, idx]-ana[0, j]) for j in range(len(times))]
    print(f'  r={rpt*100:.1f}cm 最大误差={max(errs):.3e} °C')

print('验证2: 水分扩散(常数D)')
D = 7e-9; Bi = hm*R0/D
times = [300, 600, 1200, 1800]
simC = simulate(lambda t: 0.0, lambda t: 0.05, lambda t: R0,
                lambda C, T: (1.0, 1.0, 0.0, D), 800, 0.0, C0, 1800.0, 0.5,
                1e-9, hm, save_times=times)
for rpt in [0.0, 0.01, 0.02]:
    idx = int(round(rpt/R0*800))
    ana = analytic_cyl([rpt], times, R0, C0, 0.05, D, Bi)
    errs = [abs(simC['C'][j, idx]-ana[0, j]) for j in range(len(times))]
    print(f'  r={rpt*100:.1f}cm 最大误差={max(errs):.3e} kg/kg')

# ================= 网格收敛性 =================
print('='*60)
print('网格收敛性 (问题3 烘干时长 vs 网格数N)')
def dry_time(N, dt=5.0):
    p = simulate_until(T_air_fn, C_air_fn, lambda t: R0, props_p23, N, T0, C0,
                       dt, h, hm, lambda T, C, t: np.max(C) < 0.15,
                       record_every=60.0, t_max=259200.0, n_iter=3, theta=0.5)
    return p['t_end']/3600.0
for N in [100, 200, 400, 800]:
    print(f'  N={N:4d}  烘干时长={dry_time(N):.4f} h')

print('时间步长收敛性 (问题3, N=400)')
for dt in [10.0, 5.0, 2.5, 1.0]:
    p = simulate_until(T_air_fn, C_air_fn, lambda t: R0, props_p23, 400, T0, C0,
                       dt, h, hm, lambda T, C, t: np.max(C) < 0.15,
                       record_every=60.0, t_max=259200.0, n_iter=3, theta=0.5)
    print(f'  dt={dt:5.1f}s  烘干时长={p["t_end"]/3600:.4f} h')

# ================= 参数灵敏度 =================
print('='*60)
print('问题3 烘干时长对参数的灵敏度')
base = 16.4694
def dry_time_param(props, hh=h, hmm=hm, use_R=False):
    Rf = R_fn if use_R else (lambda t: R0)
    p = simulate_until(T_air_fn, C_air_fn, Rf, props, 400, T0, C0, 5.0,
                       hh, hmm, lambda T, C, t: np.max(C) < 0.15,
                       record_every=60.0, t_max=259200.0, n_iter=3, theta=0.5)
    return p['t_end']/3600.0

print(f'  基准 h={h}, hm={hm:.1e}: {base:.4f} h')
for hh in [15.0, 20.0, 25.0, 30.0, 40.0]:
    print(f'  h={hh:5.1f}: {dry_time_param(props_p23, hh=hh):.4f} h')
for hmm in [4e-7, 6e-7, 8e-7, 1e-6, 1.2e-6]:
    print(f'  hm={hmm:.1e}: {dry_time_param(props_p23, hmm=hmm):.4f} h')

print('问题4 烘干时长对参数灵敏度 (含收缩, base=19.6583)')
for hh in [15.0, 25.0, 40.0]:
    print(f'  h={hh:5.1f}: {dry_time_param(props_p4, hh=hh, use_R=True):.4f} h')
for hmm in [4e-7, 8e-7, 1.2e-6]:
    print(f'  hm={hmm:.1e}: {dry_time_param(props_p4, hmm=hmm, use_R=True):.4f} h')

print('='*60)
print('全部完成')
