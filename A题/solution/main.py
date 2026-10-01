# -*- coding: utf-8 -*-
"""
CUMCM 2026 A题 —— 药材烘干问题 主求解脚本
生成 result1~result4.xlsx、论文表格、图表
"""
import numpy as np
import openpyxl
from openpyxl import Workbook
from solver import (simulate, simulate_until, make_air_functions, make_R_function,
                    props_p1, props_p23, props_p4, round4)

# ---------------- 全局参数 ----------------
R0 = 0.02            # 初始半径 (m)
T0 = 28.0            # 初始温度 (°C)
C0 = 2.55            # 初始水分浓度 (kg/kg)
h = 25.0             # 对流换热系数 W/(m^2·K)
hm = 8e-7            # 对流传质系数 m/s
N = 400              # 网格数（归一化半径 0..1 等分）
THETA = 0.5          # Crank-Nicolson
NITER = 3            # Picard 迭代次数

DATA = r'C:/Users/hp/Downloads/CUMCM2026Problems/A题/附件'

# ---------------- 载入数据 ----------------
def load_air():
    wb = openpyxl.load_workbook(DATA + r'/附件1.xlsx', data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))[1:]
    t = np.array([r[0] for r in rows], float)
    T = np.array([r[1] for r in rows], float)
    C = np.array([r[2] for r in rows], float)
    return t, T, C

def load_radius():
    wb = openpyxl.load_workbook(DATA + r'/附件2.xlsx', data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))[1:]
    t = np.array([r[0] for r in rows], float)
    R = np.array([r[1] for r in rows], float)
    return t, R

t_air, T_air, C_air = load_air()
t_rad, R_rad = load_radius()
T_air_fn, C_air_fn = make_air_functions(t_air, T_air, C_air, T_const=50.0, C_const=0.05)
R_fn = make_R_function(t_rad, R_rad)

# ---------------- 输出辅助 ----------------
def write_result(path, sheets, times, dists):
    """写 result 文件。sheets: list of (name, values_array); values_array 形状 (nt, nd)"""
    wb = Workbook()
    first = True
    for name, vals in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        header = ['时间\\到药材中心的距离'] + list(dists)
        ws.append(header)
        for i, tt in enumerate(times):
            row = [tt] + [round4(vals[i, j]) for j in range(len(dists))]
            ws.append(row)
    wb.save(path)

def write_result4(path, times, dists, Cvals, R_of_time):
    """result4: 固定列 0..1.9 + 药材表面，超出当前半径的格子留空"""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    header = ['时间\\到药材中心的距离'] + list(dists) + ['药材表面']
    ws.append(header)
    for i, tt in enumerate(times):
        Rt = R_of_time(i)
        row = [tt]
        for j, d in enumerate(dists):
            if d < Rt * 100.0:   # d(cm) < R(cm)
                row.append(round4(Cvals[i, j]))
            else:
                row.append(None)
        row.append(round4(Cvals[i, -1]))   # 表面值
        ws.append(row)
    wb.save(path)

def fmt_table(rows, col_labels, row_labels):
    """打印格式化表格"""
    from io import StringIO
    out = StringIO()
    out.write('          ' + ''.join(f'{c:>10}' for c in col_labels) + '\n')
    for rlab, r in zip(row_labels, rows):
        out.write(f'{rlab:>10}' + ''.join(f'{v:>10.4f}' for v in r) + '\n')
    return out.getvalue()

# ======================================================================
# 问题 1
# ======================================================================
print('#' * 70)
print('问题1：预热平衡阶段 (附录2 常数物性)')
p1 = simulate(T_air_fn, C_air_fn, lambda t: R0, props_p1, N, T0, C0, 1800.0,
              0.5, h, hm, save_times=np.arange(1, 1801), n_iter=NITER, theta=THETA)
# 表1/表2 的时刻与距离
t_tab = [100, 300, 600, 900, 1200, 1500, 1800]
d_tab = [0.0, 0.5, 1.0, 1.5, 2.0]
d_nodes = [int(round(d / 2.0 * N)) for d in d_tab]      # d(cm) -> 节点
T_tab = np.array([[p1['T'][t - 1, dn] for dn in d_nodes] for t in t_tab])
C_tab = np.array([[p1['C'][t - 1, dn] for dn in d_nodes] for t in t_tab])
print('表1 30分钟内药材的温度 (°C):')
print(fmt_table(T_tab, d_tab, [f'{t}s' for t in t_tab]))
print('表2 30分钟内药材的水分浓度 (kg/kg):')
print(fmt_table(C_tab, d_tab, [f'{t}s' for t in t_tab]))

# result1.xlsx：全时间 1..1800s，全距离 0..2.0cm
d_full = np.round(np.arange(0, 2.05, 0.1), 1)
d_nodes_full = [int(round(d / 2.0 * N)) for d in d_full]
times1 = np.arange(1, 1801)
T_out = np.array([p1['T'][i, d_nodes_full] for i in range(len(times1))])
C_out = np.array([p1['C'][i, d_nodes_full] for i in range(len(times1))])
write_result('result1.xlsx', [('温度', T_out), ('水分浓度', C_out)], times1, d_full)
print('result1.xlsx 已保存 (1800 行 x 21 列 x 2 工作表)')

# ======================================================================
# 问题 2
# ======================================================================
print('#' * 70)
print('问题2：整个烘干过程前3小时 (附录3 变物性)')
p2 = simulate(T_air_fn, C_air_fn, lambda t: R0, props_p23, N, T0, C0, 10800.0,
              0.5, h, hm, save_times=np.arange(1, 10801), n_iter=NITER, theta=THETA)
t_tab2 = [0.5, 1.0, 1.5, 2.0, 2.5, 3.0]   # 小时
T_tab2 = np.array([[p2['T'][int(tt * 3600) - 1, dn] for dn in d_nodes] for tt in t_tab2])
C_tab2 = np.array([[p2['C'][int(tt * 3600) - 1, dn] for dn in d_nodes] for tt in t_tab2])
print('表3 3小时内药材的温度 (°C):')
print(fmt_table(T_tab2, d_tab, [f'{tt}h' for tt in t_tab2]))
print('表4 3小时内药材的水分浓度 (kg/kg):')
print(fmt_table(C_tab2, d_tab, [f'{tt}h' for tt in t_tab2]))

times2 = np.arange(1, 10801)
T_out2 = np.array([p2['T'][i, d_nodes_full] for i in range(len(times2))])
C_out2 = np.array([p2['C'][i, d_nodes_full] for i in range(len(times2))])
write_result('result2.xlsx', [('温度', T_out2), ('水分浓度', C_out2)], times2, d_full)
print('result2.xlsx 已保存 (10800 行 x 21 列 x 2 工作表)')

# ======================================================================
# 问题 3
# ======================================================================
print('#' * 70)
print('问题3：确定烘干时长 (水分<0.15 kg/kg, 附录3, 无收缩)')
def stop3(T, C, t):
    return np.max(C) < 0.15

p3 = simulate_until(T_air_fn, C_air_fn, lambda t: R0, props_p23, N, T0, C0,
                    5.0, h, hm, stop3, record_every=60.0, t_max=259200.0,
                    n_iter=NITER, theta=THETA)
t_dry3 = p3['t_end']
print(f'烘干时长 t_dry3 = {t_dry3:.1f} s = {t_dry3/3600:.4f} h')
# 表5：每隔6h + 烘干结束
t_tab5 = list(range(6 * 3600, int(t_dry3), 6 * 3600))
C_tab5 = []
for tt in t_tab5:
    j = np.argmin(np.abs(p3['save_times'] - tt))   # 精确定位 60s 记录
    C_tab5.append([p3['C'][j, dn] for dn in d_nodes])
# 烘干结束时刻
j_end = np.argmin(np.abs(p3['save_times'] - t_dry3))
C_tab5.append([p3['C'][j_end, dn] for dn in d_nodes])
t_tab5_labels = [f'{tt/3600:.0f}h' for tt in t_tab5] + ['结束']
print('表5 药材烘干过程的水分浓度 (kg/kg):')
print(fmt_table(np.array(C_tab5), d_tab, t_tab5_labels))

# result3.xlsx：每 60s，距离 0..2.0cm
times3 = p3['save_times']
C_out3 = np.array([p3['C'][i, d_nodes_full] for i in range(len(times3))])
write_result('result3.xlsx', [('Sheet1', C_out3)], times3, d_full)
print(f'result3.xlsx 已保存 ({len(times3)} 行 x 21 列)')

# ======================================================================
# 问题 4
# ======================================================================
print('#' * 70)
print('问题4：考虑收缩 (附录4 变物性 + 半径变化)')
def stop4(T, C, t):
    return np.max(C) < 0.15

p4 = simulate_until(T_air_fn, C_air_fn, R_fn, props_p4, N, T0, C0,
                    5.0, h, hm, stop4, record_every=60.0, t_max=259200.0,
                    n_iter=NITER, theta=THETA)
t_dry4 = p4['t_end']
print(f'烘干时长 t_dry4 = {t_dry4:.1f} s = {t_dry4/3600:.4f} h')
# 表6：距离 0,0.5,1.0,1.5 + 药材表面
d_tab6 = [0.0, 0.5, 1.0, 1.5]
t_tab6 = list(range(6 * 3600, int(t_dry4), 6 * 3600))
rows6 = []
labels6 = []
for tt in t_tab6:
    j = np.argmin(np.abs(p4['save_times'] - tt))
    Rt = p4['R'][j]
    row = []
    for d in d_tab6:
        dm = d / 100.0
        if dm < Rt:
            xi_d = dm / Rt
            row.append(np.interp(xi_d, p4['xi'], p4['C'][j]))
        else:
            row.append(np.nan)
    row.append(p4['C'][j, -1])   # 表面
    rows6.append(row)
    labels6.append(f'{tt/3600:.0f}h')
j_end = np.argmin(np.abs(p4['save_times'] - t_dry4))
Rt = p4['R'][j_end]
row = []
for d in d_tab6:
    dm = d / 100.0
    if dm < Rt:
        xi_d = dm / Rt
        row.append(np.interp(xi_d, p4['xi'], p4['C'][j_end]))
    else:
        row.append(np.nan)
row.append(p4['C'][j_end, -1])
rows6.append(row)
labels6.append('结束')
print('表6 药材烘干过程的水分浓度 (kg/kg) (列: 0,0.5,1.0,1.5,表面):')
print(fmt_table(np.array(rows6), d_tab6 + ['表面'], labels6))

# result4.xlsx：每 60s，距离 0..1.9 + 药材表面
d_full4 = np.round(np.arange(0, 1.95, 0.1), 1)
times4 = p4['save_times']
# 预计算每个时刻、每个固定距离上的 C（超出半径的用 nan 占位）
C_out4_fixed = np.full((len(times4), len(d_full4)), np.nan)
for i in range(len(times4)):
    Rt = p4['R'][i]
    for j, d in enumerate(d_full4):
        dm = d / 100.0
        if dm < Rt:
            xi_d = dm / Rt
            C_out4_fixed[i, j] = np.interp(xi_d, p4['xi'], p4['C'][i])
C_surf4 = p4['C'][:, -1]
# 写 result4
wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'
ws.append(['时间\\到药材中心的距离'] + list(d_full4) + ['药材表面'])
for i, tt in enumerate(times4):
    row = [tt]
    for j in range(len(d_full4)):
        if not np.isnan(C_out4_fixed[i, j]):
            row.append(round4(C_out4_fixed[i, j]))
        else:
            row.append(None)
    row.append(round4(C_surf4[i]))
    ws.append(row)
wb.save('result4.xlsx')
print(f'result4.xlsx 已保存 ({len(times4)} 行 x {len(d_full4)+1} 列)')

# ======================================================================
# 保存关键中间量供绘图/论文使用
# ======================================================================
np.savez('_results.npz',
         p1_T=p1['T'], p1_C=p1['C'], p1_times=p1['save_times'], p1_xi=p1['xi'],
         p2_T=p2['T'], p2_C=p2['C'], p2_times=p2['save_times'], p2_xi=p2['xi'],
         p3_C=p3['C'], p3_times=p3['save_times'], p3_xi=p3['xi'], t_dry3=t_dry3,
         p4_C=p4['C'], p4_R=p4['R'], p4_times=p4['save_times'], p4_xi=p4['xi'], t_dry4=t_dry4,
         t_air=t_air, T_air=T_air, C_air=C_air, t_rad=t_rad, R_rad=R_rad)
print('_results.npz 已保存')
print('#' * 70)
print('全部问题求解完成！')
